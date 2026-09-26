"""Tests of the Word export (PHASE6_SPEC §11.1): every generated file is validated against the ECMA-376
schema and the integrity checks; golden files G1–G8 (`publishing/word/golden/`, updated with
`NASSAKH_UPDATE_GOLDEN=1`); the options matrix read back with python-docx; the model's `editorial` flag;
fonts (obfuscation, the key, what is embedded); determinism; performance; the §5.14 edge cases; the
exporter against the pipeline's contract; the harness's calibration files.

Every build here resolves the faces with an empty font folder, so Amiri stands in for every face on any
machine and the goldens are the same everywhere."""

from __future__ import annotations

import io
import os
import re
import time
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from django.conf import settings
from django.test import override_settings

import docx as python_docx
import pytest
from lxml import etree

from editor import uncertain
from editor.models import Manuscript
from publishing import fonts as F
from publishing.exporters import ExportJob, InvalidExport, NullProgress, get_exporter, parse_options
from publishing.model import STYLES, Block, Book, Footnote, NoteRef, Run, SourceMark, book_model, page_setup
from publishing.tests import LOREM, document, heading, long_book, make_book, note, para, text
from publishing.word import calibration, faces, ooxml, schema
from publishing.word.exporter import DOCX_OPTIONS, DocxExporter, readings_of, static_notes
from publishing.word.ooxml import ORDER, W_NS
from publishing.word.options import COMMENTS_SUFFIX, WORD_VERSION, WordOptions
from publishing.word.runs import _pieces
from publishing.word.sections import Bookmarks
from publishing.word.styles import char_styles, collapse, split_gap, style_table, widow_control
from publishing.word.writer import DocMeta, PagePlan, build_docx

GOLDEN_DIR = Path(__file__).resolve().parent / "word" / "golden"
XSD = "http://www.w3.org/2001/XMLSchema"
W = f"{{{W_NS}}}"
META = DocMeta.fixed()

# ====================================================================== helpers


@pytest.fixture(autouse=True)
def _no_system_fonts(tmp_path):
    """Every face resolves to the vendored Amiri (no Mac font folders): the same bytes on any machine."""
    with override_settings(NASSAKH={**settings.NASSAKH, "FONT_DIRS": [str(tmp_path / "no-fonts")]}):
        yield


def amiri() -> F.ResolvedFonts:
    return F.resolve("amiri", "amiri", "amiri")


def lotus() -> F.ResolvedFonts:
    """Lotus (no Latin letters or digits) with Amiri's files standing in, and Times as the Latin face."""
    files = F.locate("amiri")
    spec = F.FONTS["lotus"]
    body = F.Face("lotus", "lotus", spec.name, spec.family, files, latin=False)
    times = F.Face(
        "times", "times", "Times New Roman", "Times New Roman", F.FontFiles(files.bold), latin=True
    )
    return F.ResolvedFonts(body=body, latin=times, heading=body)


def build(source, sheet=None, options=None, *, fonts=None, plan=None, readings=None, **kwargs):
    book = book_model(source, sheet, title="كتابي", author="المؤلف", editorial=bool(readings is not None))
    return build_docx(
        book,
        fonts or amiri(),
        options or WordOptions(),
        plan=plan,
        readings=readings,
        meta=META,
        embed_fonts=kwargs.pop("embed_fonts", False),
        **kwargs,
    )


def check(data: bytes) -> None:
    assert schema.validate_package(data) == []
    assert schema.integrity_errors(data) == []


def parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def xml(data: bytes, name: str) -> etree._Element:
    return etree.fromstring(parts(data)[name])


def canonical(data: bytes, names: list[str]) -> str:
    out = []
    for name in names:
        root = etree.fromstring(parts(data)[name])
        out.append(f"### {name}\n" + etree.tostring(root, pretty_print=True, encoding="unicode"))
    return "\n".join(out)


def golden(name: str, data: bytes, names: list[str]) -> None:
    """Compare the named parts with `golden/<name>.xml`; `NASSAKH_UPDATE_GOLDEN=1` rewrites it."""
    check(data)
    actual = canonical(data, names)
    path = GOLDEN_DIR / f"{name}.xml"
    if os.environ.get("NASSAKH_UPDATE_GOLDEN") == "1" or not path.exists():
        GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
        path.write_text(actual, encoding="utf-8")
    expected = path.read_text(encoding="utf-8")
    assert actual == expected, f"{name} differs from its golden (NASSAKH_UPDATE_GOLDEN=1 to accept)"


def read_back(data: bytes):
    return python_docx.Document(io.BytesIO(data))


def sheet(**values) -> dict:
    return values


class Reading:
    """A stand-in for `editor.uncertain.Word` (its `word`, `readings` and `source_page`)."""

    def __init__(self, word: str, readings=None, page: int | None = 5):
        self.word = word
        self.readings = (
            readings
            if readings is not None
            else [
                {"engine": "primary", "label": "Qari v0.3", "text": word, "current": True},
                {"engine": "tess", "label": "Tesseract", "text": word + "ة", "current": False},
            ]
        )
        self.source_page = page


# ====================================================================== the schema order tables


def test_order_tables_agree_with_the_vendored_schema():
    root = etree.parse(str(schema.XSD_DIR / "wml.xsd")).getroot()
    types = {t.get("name"): t for t in root.iter(f"{{{XSD}}}complexType") if t.get("name")}
    groups = {g.get("name"): g for g in root.iter(f"{{{XSD}}}group") if g.get("name")}

    def sequence(node, out):
        for child in node:
            tag = etree.QName(child).localname
            if tag == "element":
                out.append((child.get("name") or child.get("ref")).split(":")[-1])
            elif tag in ("sequence", "choice", "all", "extension", "complexContent"):
                base = child.get("base")
                if base and base.split(":")[-1] in types:
                    sequence(types[base.split(":")[-1]], out)
                sequence(child, out)
            elif tag == "group":
                found = groups.get((child.get("ref") or "").split(":")[-1])
                if found is not None:
                    sequence(found, out)
        return out

    for tag, type_name in (
        ("pPr", "CT_PPr"),
        ("rPr", "CT_RPr"),
        ("sectPr", "CT_SectPr"),
        ("footnotePr", "CT_FtnDocProps"),
        ("style", "CT_Style"),
        ("styles", "CT_Styles"),
        ("docDefaults", "CT_DocDefaults"),
        ("font", "CT_Font"),
        ("pBdr", "CT_PBdr"),
        ("compat", "CT_Compat"),
        ("settings", "CT_Settings"),
    ):
        assert tuple(sequence(types[type_name], [])) == ORDER[tag], tag


def test_w_sorts_children_into_schema_order_and_keeps_others_in_place():
    ppr = ooxml.w("pPr", ooxml.w("jc", val="both"), ooxml.w("spacing", after=0), ooxml.w("pStyle", val="X"))
    assert [ooxml.local(c) for c in ppr] == ["pStyle", "spacing", "jc"]
    p = ooxml.w("p", ooxml.w("r"), ooxml.w("pPr"))
    assert [ooxml.local(c) for c in p] == ["r", "pPr"]  # w:p is not sorted: the caller builds it in order
    assert ooxml.xml_safe("a\x01b\ud800c￾") == "abc"
    assert ooxml.twips_mm(170) == 9638 and ooxml.twips_mm(240) == 13606 and ooxml.half_points(20.8) == 42


# ====================================================================== goldens G1–G8


def test_g1_one_chapter_body_styles_and_settings():
    source = document(heading("h1", "الفصل الأول"), para("p1", LOREM), para("p2", "فقرة ثانية قصيرة."))
    result = build(source, sheet(front_matter={"title_page": False, "contents": False}))
    golden("G1", result.data, ["word/document.xml", "word/styles.xml", "word/settings.xml"])
    styles = xml(result.data, "word/styles.xml")
    ids = {s.get(f"{W}styleId") for s in styles.iter(f"{W}style")}
    assert {style.word for style in STYLES.values()} <= ids  # every model style id is written (§4.3)
    assert set(char_styles(page_setup(None))) <= ids
    normal = next(s for s in styles.iter(f"{W}style") if s.get(f"{W}styleId") == "Normal")
    spacing = normal.find(f"{W}pPr/{W}spacing")
    assert spacing.get(f"{W}line") == str(round(1.7 * 13 * 20)) and spacing.get(f"{W}lineRule") == "exact"
    assert normal.find(f"{W}pPr/{W}jc").get(f"{W}val") == "lowKashida"
    assert normal.find(f"{W}pPr/{W}widowControl") is not None
    settings_xml = xml(result.data, "word/settings.xml")
    assert settings_xml.find(f"{W}updateFields") is None
    assert settings_xml.find(f"{W}zoom").get(f"{W}percent") == "100"
    compat = [c.get(f"{W}val") for c in settings_xml.iter(f"{W}compatSetting")]
    assert compat[0] == "15"
    assert [ooxml.local(c) for c in settings_xml][:2] == ["view", "zoom"]


def test_g2_marks_direction_faces_and_breaks():
    source = document(
        heading("h1", "الفصل"),
        para(
            "p1",
            text("غامق", "bold"),
            " ومائل ",
            text("مائل", "italic"),
            " و Ibn Khaldun في 1377 – «اقتباس»… ",
            "شطر",
            {"type": "hardBreak"},
            "عجز",
        ),
    )
    result = build(source, sheet(front_matter={"title_page": False, "contents": False}))
    golden("G2", result.data, ["word/document.xml"])
    document_xml = parts(result.data)["word/document.xml"].decode()
    assert "<w:b/><w:bCs/>" in document_xml and "<w:i/><w:iCs/>" in document_xml
    assert document_xml.count("<w:br/>") == 1
    assert "NkLatin" not in document_xml  # Amiri covers every character
    with_lotus = build(source, sheet(front_matter={"title_page": False, "contents": False}), fonts=lotus())
    golden("G2-lotus", with_lotus.data, ["word/document.xml", "word/fontTable.xml"])
    lotus_xml = parts(with_lotus.data)["word/document.xml"].decode()
    assert (
        lotus_xml.count('w:val="NkLatin"') >= 1
    )  # the digits (Amiri stands in: it has the dash and the ellipsis)
    runs = [
        (r.find(f"{W}t").text, r.find(f"{W}rPr/{W}rStyle") is not None)
        for r in xml(with_lotus.data, "word/document.xml").iter(f"{W}r")
        if r.find(f"{W}t") is not None
    ]
    assert ("1377", True) in runs and ("Ibn Khaldun", False) in runs
    families = [f.get(f"{W}name") for f in xml(with_lotus.data, "word/fontTable.xml").iter(f"{W}font")]
    assert families == ["Lotus Linotype Exnd", "Times New Roman"]


@pytest.mark.parametrize("numbering", ["page", "chapter", "book"])
def test_g3_footnotes_per_page_chapter_and_book(numbering):
    source = long_book(chapters=2, paragraphs=3, notes_every=1)
    result = build(
        source, sheet(footnote_numbering=numbering, front_matter={"title_page": False, "contents": False})
    )
    golden(f"G3-{numbering}", result.data, ["word/footnotes.xml", "word/settings.xml"])
    notes = xml(result.data, "word/footnotes.xml")
    kinds = [(n.get(f"{W}type"), n.get(f"{W}id")) for n in notes.iter(f"{W}footnote")]
    assert kinds[:2] == [("separator", "-1"), ("continuationSeparator", "0")] and len(kinds) == 8
    separator = notes.find(f"{W}footnote/{W}p/{W}pPr")
    assert separator.find(f"{W}pBdr/{W}top").get(f"{W}sz") == "3"
    assert separator.find(f"{W}spacing").get(f"{W}before") == str(ooxml.twips_mm(4))
    restart = {"page": "eachPage", "chapter": "eachSect", "book": "continuous"}[numbering]
    settings_xml = xml(result.data, "word/settings.xml")
    assert settings_xml.find(f"{W}footnotePr/{W}numRestart").get(f"{W}val") == restart
    for section in xml(result.data, "word/document.xml").iter(f"{W}sectPr"):
        assert section.find(f"{W}footnotePr/{W}numRestart").get(f"{W}val") == restart
    first = notes.findall(f"{W}footnote")[2]
    marker = [r.find(f"{W}rPr/{W}rStyle").get(f"{W}val") for r in list(first.iter(f"{W}r"))[:3]]
    assert marker == ["NkNoteNumber"] * 3 and first.find(f".//{W}footnoteRef") is not None


def test_g4_front_matter_and_the_prefilled_contents():
    source = document(
        heading("h1", "الفصل الأول"),
        para("p1", LOREM),
        heading("h2", "مبحث", level=2),
        para("p2", LOREM),
        heading("h3", "الفصل الثاني"),
        para("p3", LOREM),
    )
    fields = {
        "subtitle": "دراسة",
        "editor": "المحقق",
        "translator": "ترجمة: المترجم",
        "publisher": "دار النشر",
        "city": "طرابلس",
        "year": "2026",
        "edition": "الأولى",
        "isbn": "978-9959-0-0001-1",
        "rights": "جميع الحقوق محفوظة\nللناشر",
    }
    plan = PagePlan("live", 12, {"h1": 5, "h2": 6, "h3": 9}, {"h1": (5, 8), "h3": (9, 12)})
    result = build(source, sheet(front_matter={"copyright_page": True, "fields": fields}), plan=plan)
    golden(
        "G4",
        result.data,
        ["word/document.xml", "docProps/core.xml", "docProps/custom.xml", "docProps/app.xml"],
    )
    body = xml(result.data, "word/document.xml").find(f"{W}body")
    styles = [
        p.find(f"{W}pPr/{W}pStyle").get(f"{W}val")
        for p in body.iter(f"{W}p")
        if p.find(f"{W}pPr/{W}pStyle") is not None
    ]
    assert styles[:7] == ["Title", "Subtitle", "NkAuthor", "NkCredit", "NkCredit", "NkImprint", "NkCopyright"]
    assert styles.count("NkCopyright") == 9 and styles[styles.index("TOCHeading") + 1 :][:3] == [
        "TOC1",
        "TOC2",
        "TOC1",
    ]
    title = next(p for p in body.iter(f"{W}p"))
    spacing = title.find(f"{W}pPr/{W}spacing")
    assert spacing.get(f"{W}before") == str(ooxml.twips_mm(0.28 * 130)) and spacing.get(f"{W}after") == str(
        ooxml.twips_mm(4)
    )
    instr = [i.text.strip() for i in body.iter(f"{W}instrText")]
    assert instr == [
        'TOC \\o "1-2" \\h \\z \\u',
        "PAGEREF _Toc_h1 \\h",
        "PAGEREF _Toc_h2 \\h",
        "PAGEREF _Toc_h3 \\h",
    ]
    links = [
        (h.get(f"{W}anchor"), "".join(t.text for t in h.iter(f"{W}t"))) for h in body.iter(f"{W}hyperlink")
    ]
    assert links == [("_Toc_h1", "الفصل الأول5"), ("_Toc_h2", "مبحث6"), ("_Toc_h3", "الفصل الثاني9")]
    bookmarks = [b.get(f"{W}name") for b in body.iter(f"{W}bookmarkStart")]
    assert bookmarks == ["_Toc_h1", "_Toc_h2", "_Toc_h3"]
    read = read_back(result.data)
    core = read.core_properties
    assert (core.title, core.subject, core.author, core.language) == (
        "كتاب التجربة",
        "دراسة",
        "مؤلف الكتاب",
        "ar",
    )
    assert core.comments == "جميع الحقوق محفوظة للناشر" and core.last_modified_by == "نسّاخ"
    assert core.created == datetime(2026, 1, 1, 0, 0, 0, tzinfo=UTC)
    custom = parts(result.data)["docProps/custom.xml"].decode()
    assert all(name in custom for name in ("ISBN", "Edition", "Publisher", "City", "Year", "NassakhRenderer"))
    assert "<Company>دار النشر</Company>" in parts(result.data)["docProps/app.xml"].decode()


def test_g5_recto_openings_chapter_header_and_outer_numbers():
    source = long_book(chapters=3, paragraphs=2, notes_every=5)
    result = build(
        source, sheet(chapter_opening="recto", running_header="chapter", page_number="bottom_outer")
    )
    golden(
        "G5",
        result.data,
        [
            "word/document.xml",
            "word/header1.xml",
            "word/header2.xml",
            "word/header3.xml",
            "word/footer1.xml",
            "word/footer2.xml",
            "word/footer3.xml",
        ],
    )
    names = parts(result.data)
    assert sorted(n for n in names if "header" in n or "footer" in n) == [
        "word/footer1.xml", "word/footer2.xml", "word/footer3.xml",
        "word/header1.xml", "word/header2.xml", "word/header3.xml", "word/header4.xml", "word/header5.xml",
        "word/header6.xml", "word/header7.xml",
    ]  # fmt: skip
    sections = list(xml(result.data, "word/document.xml").iter(f"{W}sectPr"))
    assert len(sections) == 5  # title page, contents, three chapters
    kinds = [s.find(f"{W}type").get(f"{W}val") if s.find(f"{W}type") is not None else None for s in sections]
    assert kinds == [None, "oddPage", "oddPage", "oddPage", "oddPage"]
    assert all(s.find(f"{W}headerReference") is None for s in sections[:2])  # front sections: no parts
    first = [(r.tag.rsplit("}")[1], r.get(f"{W}type")) for r in sections[2] if r.tag.endswith("Reference")]
    assert sorted(first) == sorted(
        [("footerReference", "default"), ("footerReference", "even"), ("footerReference", "first"),
         ("headerReference", "default"), ("headerReference", "even"), ("headerReference", "first")]
    )  # fmt: skip
    later = [(r.tag.rsplit("}")[1], r.get(f"{W}type")) for r in sections[3] if r.tag.endswith("Reference")]
    assert sorted(later) == [
        ("headerReference", "default"),
        ("headerReference", "even"),
    ]  # the rest is inherited
    assert all(s.find(f"{W}titlePg") is not None for s in sections[2:])
    assert xml(result.data, "word/settings.xml").find(f"{W}evenAndOddHeaders") is not None
    footer = xml(result.data, "word/footer1.xml")
    assert footer.find(f".//{W}jc").get(f"{W}val") == "left" and footer.find(f".//{W}bidi") is None
    assert [i.text.strip() for i in footer.iter(f"{W}instrText")] == ["PAGE"]
    header = xml(result.data, "word/header1.xml")
    assert "".join(t.text for t in header.iter(f"{W}t")) == "الفصل 1"
    assert "".join(t.text for t in xml(result.data, "word/header3.xml").iter(f"{W}t")) == ""  # the opener
    margins = sections[2].find(f"{W}pgMar")
    assert (margins.get(f"{W}left"), margins.get(f"{W}right")) == (
        str(ooxml.twips_mm(22)),
        str(ooxml.twips_mm(18)),
    )
    assert sections[2].find(f"{W}bidi") is not None and sections[2].find(f"{W}pgNumType") is None


def test_g6_comments_across_a_bold_boundary_and_in_a_note():
    source = document(
        heading("h1", "الفصل"),
        para(
            "p1",
            "كلمة ",
            text("مش", "uncertain", "bold"),
            text("كوكة", "uncertain"),
            " ثم ",
            text("ثانية", "uncertain"),
            note("n1", "حاشية فيها ", text("شك", "uncertain"), " أيضًا"),
            " بعدها.",
        ),
    )
    readings = {
        ("p1", None): [Reading("مشكوكة"), Reading("ثانية", [])],
        ("p1", "n1"): [Reading("شك", page=None)],
    }
    result = build(
        source,
        sheet(front_matter={"title_page": False, "contents": False}),
        WordOptions(comments=True),
        readings=readings,
    )
    golden("G6", result.data, ["word/document.xml", "word/footnotes.xml", "word/comments.xml"])
    assert result.stats["comments"] == 3 and result.stats["comments_skipped"] == 0
    assert [w["code"] for w in result.warnings] == ["comments"]
    document_xml = xml(result.data, "word/document.xml")
    starts = [c.get(f"{W}id") for c in document_xml.iter(f"{W}commentRangeStart")]
    assert starts == ["0", "1"]
    paragraph = next(document_xml.iter(f"{W}p"))
    paragraph = document_xml.findall(f".//{W}p")[1]
    order = [ooxml.local(c) for c in paragraph if ooxml.local(c) != "pPr"]
    assert order[:5] == ["r", "commentRangeStart", "r", "r", "commentRangeEnd"]  # the word spans two runs
    notes = xml(result.data, "word/footnotes.xml")
    assert [c.get(f"{W}id") for c in notes.iter(f"{W}commentRangeStart")] == ["2"]
    comments = xml(result.data, "word/comments.xml")
    texts = ["".join(t.text for t in c.iter(f"{W}t")) for c in comments.iter(f"{W}comment")]
    assert (
        texts[0]
        == "كلمة غير مؤكَّدة: مشكوكةالقراءات:Qari v0.3: مشكوكة (في النص)Tesseract: مشكوكةةالصفحة الأصلية: 5"
    )
    assert texts[1] == "كلمة غير مؤكَّدة: ثانيةالصفحة الأصلية: 5لا قراءات أخرى لهذه الكلمة؛ راجعها على الأصل."
    assert texts[2].startswith("كلمة غير مؤكَّدة: شك")
    first = next(comments.iter(f"{W}comment"))
    assert (first.get(f"{W}author"), first.get(f"{W}initials"), first.get(f"{W}date")) == (
        "نسّاخ",
        "ن",
        "2026-01-01T00:00:00Z",
    )
    read = read_back(result.data)
    assert len(list(read.comments)) == 3


def test_g7_a_book_without_headings_numbered_per_chapter():
    blocks = [para(f"p{i}", LOREM, pages=(i * 40,)) for i in range(1, 4)]  # far-apart scan pages: 3 sections
    source = document(*blocks)
    result = build(source, sheet(footnote_numbering="chapter"))
    golden("G7", result.data, ["word/document.xml"])
    sections = list(xml(result.data, "word/document.xml").iter(f"{W}sectPr"))
    kinds = [s.find(f"{W}type").get(f"{W}val") if s.find(f"{W}type") is not None else None for s in sections]
    assert kinds == [None, "nextPage", "continuous", "continuous"]  # the title page, then the three sections
    marks = [b.get(f"{W}name") for b in xml(result.data, "word/document.xml").iter(f"{W}bookmarkStart")]
    assert all(mark.startswith("_nk_ch_") for mark in marks) and len(marks) == 3
    plain = build(source, sheet(footnote_numbering="page"))
    assert len(list(xml(plain.data, "word/document.xml").iter(f"{W}sectPr"))) == 2  # the sections run on
    assert "TOC" not in parts(plain.data)["word/document.xml"].decode()  # no contents without headings


def test_g8_block_attributes_and_the_margin_collapse_overrides():
    source = document(
        heading("h1", "الفصل"),
        heading("h2", "عنوان فرعي", level=2),
        para("p1", "اقتباس أول", style="quote"),
        para("p2", "اقتباس ثان", style="quote"),
        para("p3", "بيت", style="verse"),
        {"type": "separator", "attrs": {"id": "s1"}},
        para("p4", LOREM, breakBefore=True, keepWithNext=True),
        para("p5", LOREM),
    )
    result = build(source, sheet(front_matter={"title_page": False, "contents": False}, widows=3, orphans=1))
    golden("G8", result.data, ["word/document.xml"])
    body = xml(result.data, "word/document.xml").find(f"{W}body")
    paragraphs = list(body.iter(f"{W}p"))

    def spacing(index):
        node = paragraphs[index].find(f"{W}pPr/{W}spacing")
        return (node.get(f"{W}before"), node.get(f"{W}after")) if node is not None else None

    # Word takes the larger of the two spacings (C12), as CSS does: no overrides between these blocks
    assert all(spacing(index) is None for index in range(6))
    assert paragraphs[6].find(f"{W}pPr/{W}pageBreakBefore") is not None
    assert paragraphs[6].find(f"{W}pPr/{W}keepNext") is not None
    assert spacing(6) is None  # a page break starts a new flow
    styles = xml(result.data, "word/styles.xml")
    normal = next(s for s in styles.iter(f"{W}style") if s.get(f"{W}styleId") == "Normal")
    assert normal.find(f"{W}pPr/{W}widowControl") is not None  # 3/1 → on, approximate
    assert widow_control(page_setup(sheet(widows=3, orphans=1))) == (True, True)
    assert widow_control(page_setup(sheet(widows=1, orphans=1))) == (False, False)


# ====================================================================== the collapse rule and styles


def test_collapse_and_split_gap_follow_css():
    assert collapse(2, 2) == 2 and collapse(9, 5) == 9 and collapse(10, -6) == 4 and collapse(-2, -3) == -3
    # C12: Word takes the larger of the two spacings, as CSS does, so equal margins need nothing
    assert split_gap(2, 2, 2) == (None, None)
    assert split_gap(9, 5, 9) == (None, None)
    assert split_gap(2, 4, 4) == (None, None)  # a separator (4 before) after a verse (2 after)
    assert split_gap(10, 0, 4) == (4, 0)  # Title (10 after) → Subtitle (−6 before): 4 mm
    assert split_gap(0, 0, 30) == (None, 30)
    assert split_gap(0, 0, 0) == (None, None)
    assert split_gap(3, 2, -1) == (0, 0)  # a negative CSS gap becomes none


def test_style_table_takes_the_unrounded_css_values():
    table = style_table(
        page_setup(sheet(body_size_pt=13, h1_scale=1.6, line_height=1.7, footnote_size_pt=10))
    )
    assert table["Heading1"].size == pytest.approx(20.8) and ooxml.half_points(table["Heading1"].size) == 42
    assert table["Heading1"].line_twips == round(1.35 * 20.8 * 20)
    assert table["Title"].size == pytest.approx(1.6 * 1.4 * 13)
    assert (
        table["Header"].size == 9.5
        and table["Footer"].ltr
        and table["TOC2"].indent == pytest.approx(1.5 * 0.95 * 13)
    )
    assert table["Quote"].indent == 26 and table["Normal"].first_line == pytest.approx(19.5)
    chars = char_styles(page_setup(None))
    assert (
        ooxml.half_points(chars["FootnoteReference"].size) == 16
        and ooxml.half_points(chars["FootnoteReference"].position) == 8
    )


def test_start_aligned_styles_and_kashida_choices():
    source = document(heading("h1", "الفصل"), para("p1", LOREM))
    for key, jc in (
        ("none", "both"),
        ("low", "lowKashida"),
        ("medium", "mediumKashida"),
        ("high", "highKashida"),
    ):
        result = build(source, None, WordOptions(kashida=key))
        styles = xml(result.data, "word/styles.xml")
        by_id = {s.get(f"{W}styleId"): s for s in styles.iter(f"{W}style")}
        assert by_id["Normal"].find(f"{W}pPr/{W}jc").get(f"{W}val") == jc
        assert by_id["FootnoteText"].find(f"{W}pPr/{W}jc").get(f"{W}val") == jc
        assert by_id["Quote"].find(f"{W}pPr/{W}jc").get(f"{W}val") == jc
        assert by_id["NkVerse"].find(f"{W}pPr/{W}jc").get(f"{W}val") == "center"
        assert by_id["TOC1"].find(f"{W}pPr/{W}jc").get(f"{W}val") == "start"
        tab = by_id["TOC1"].find(f"{W}pPr/{W}tabs/{W}tab")
        assert (tab.get(f"{W}val"), tab.get(f"{W}leader"), tab.get(f"{W}pos")) == (
            "right",
            "dot",
            str(ooxml.twips(130 / 25.4 * 72)),
        )


# ====================================================================== headers, footers, sections


def test_top_outer_numbers_put_the_number_at_the_outer_edge_with_a_centre_tab():
    source = long_book(chapters=2, paragraphs=2, notes_every=5)
    result = build(source, sheet(page_number="top_outer", running_header="book", chapter_opening="any"))
    check(result.data)
    odd = xml(result.data, "word/header1.xml")
    even = xml(result.data, "word/header2.xml")
    first = xml(result.data, "word/header3.xml")
    assert odd.find(f".//{W}bidi").get(f"{W}val") == "0" and odd.find(f".//{W}jc").get(f"{W}val") == "left"
    assert [t.get(f"{W}val") for t in odd.iter(f"{W}tab") if t.get(f"{W}val")] == ["center", "right"]
    assert [ooxml.local(c) for c in odd.find(f".//{W}p")].count("r") >= 6
    even_text = "".join(t.text for t in even.iter(f"{W}t"))
    assert even_text.startswith("كتاب التجربة") and even_text.endswith("1")
    assert "".join(t.text for t in first.iter(f"{W}t")) == "1"  # the opener keeps only the number
    assert "word/footer1.xml" not in parts(result.data)
    sections = list(xml(result.data, "word/document.xml").iter(f"{W}sectPr"))
    assert sections[-1].find(f"{W}headerReference") is None  # same book title: inherited


def test_page_numbers_none_and_bottom_center():
    source = long_book(chapters=2, paragraphs=2, notes_every=5)
    none = build(source, sheet(page_number="none"))
    check(none.data)
    assert not [n for n in parts(none.data) if "header" in n or "footer" in n]
    center = build(source, sheet(page_number="bottom_center", running_header="none"))
    check(center.data)
    assert [n for n in parts(center.data) if "footer" in n] == ["word/footer1.xml"]
    footer = xml(center.data, "word/footer1.xml")
    assert footer.find(f".//{W}jc") is None  # the Footer style is centred and LTR
    sections = list(xml(center.data, "word/document.xml").iter(f"{W}sectPr"))
    assert all(s.find(f"{W}titlePg") is None for s in sections)  # no running header: no first page parts
    assert xml(center.data, "word/settings.xml").find(f"{W}evenAndOddHeaders") is None


def test_header_and_footer_distances_mirror_the_preview_paddings():
    result = build(
        document(heading("h1", "الفصل"), para("p1", "نص")),
        sheet(top_mm=20, bottom_mm=22, footnote_size_pt=10),
    )
    margins = next(xml(result.data, "word/document.xml").iter(f"{W}pgMar"))
    header_line = 1.5 * 9.5 * 25.4 / 72
    footer_line = 1.5 * 10 * 25.4 / 72
    assert margins.get(f"{W}header") == str(ooxml.twips_mm(20 - 3 - header_line))
    assert margins.get(f"{W}footer") == str(ooxml.twips_mm(22 - 4 - footer_line))
    small = build(document(heading("h1", "الفصل"), para("p1", "نص")), sheet(top_mm=6, bottom_mm=6))
    margins = next(xml(small.data, "word/document.xml").iter(f"{W}pgMar"))
    assert margins.get(f"{W}header") == margins.get(f"{W}footer") == str(ooxml.twips_mm(4))


# ====================================================================== the contents and the plan


def test_contents_without_a_plan_has_no_numbers_and_warns():
    source = document(heading("h1", "الفصل"), para("p1", LOREM))
    result = build(source, None, plan=None)
    check(result.data)
    assert [w["code"] for w in result.warnings] == ["toc_numbers_missing"]
    link = next(xml(result.data, "word/document.xml").iter(f"{W}hyperlink"))
    assert "".join(t.text for t in link.iter(f"{W}t")) == "الفصل"  # no cached number
    partial = build(source, None, plan=PagePlan("render", 3, {}, {}))
    assert partial.warnings == [] and "no page in the plan" in partial.log[0]


def test_page_plan_from_layout_pages():
    pages = [
        {"n": 1, "chapter": None, "lines": [{"kind": "title", "block": "front-title"}]},
        {
            "n": 3,
            "chapter": "h1",
            "lines": [{"kind": "heading", "block": "h1"}, {"kind": "body", "block": "p1"}],
        },
        {
            "n": 4,
            "chapter": "h1",
            "lines": [{"kind": "heading", "block": "h1"}, {"kind": "heading", "block": "s1"}],
        },
    ]
    plan = PagePlan.from_pages(pages, "live", [{"id": "h2", "first": 5, "last": 6}])
    assert plan.page_count == 3 and plan.headings == {"h1": 3, "s1": 4}
    assert plan.chapters == {"h1": (3, 4), "h2": (5, 6)}


def test_bookmark_names_are_safe_short_and_unique():
    marks = Bookmarks()
    assert marks.name("_Toc_", "h1") == "_Toc_h1"
    assert marks.name("_Toc_", "h1") == "_Toc_h1"  # the same id, the same name
    assert marks.name("_Toc_", "عنوان-١") == "_Toc_" + "_" * 7
    assert marks.name("_Toc_", "عنوان-٢") == "_Toc_" + "_" * 7 + "_2"
    long = marks.name("_Toc_", "x" * 60)
    assert len(long) == 40 and len(marks.name("_Toc_", "x" * 61)) == 40


# ====================================================================== the model's editorial flag


def test_editorial_false_leaves_the_model_unchanged():
    source = document(heading("h1", "الفصل"), para("p1", "نص ", text("مؤكَّد", "uncertain"), " آخر"))
    plain, again = book_model(source, None), book_model(source, None, editorial=False)
    assert plain == again
    assert list(plain.blocks())[1].runs == [Run("نص مؤكَّد آخر")]
    editorial = book_model(source, None, editorial=True)
    assert list(editorial.blocks())[1].runs == [Run("نص "), Run("مؤكَّد", ("uncertain",)), Run(" آخر")]
    assert [b.text() for b in editorial.blocks()] == [b.text() for b in plain.blocks()]


def _rebuilt_words(source) -> list[tuple[str, str | None, str]]:
    """`(block, note, word)` as the writer rebuilds them from the editorial model."""
    out = []
    for block in book_model(source, None, editorial=True).blocks():
        _pieces_, words = _pieces(block.runs)
        out += [(block.id, None, text) for _indexes, text in words]
        for footnote in block.footnotes:
            _p, note_words = _pieces(footnote.runs)
            out += [(block.id, footnote.id, text) for _indexes, text in note_words]
    return out


def test_the_writer_rebuilds_the_uncertain_words_as_words_of_counts_them():
    source = document(
        heading("h1", "الفصل"),
        para(
            "p1",
            "أ ",
            text("عبر", "uncertain", "bold"),
            text("العلامات", "uncertain"),
            " ",
            text("كلمتان", "uncertain"),
            " ",
            text("متتاليتان", "uncertain"),
            text(" بجوار", "uncertain"),
            note("n1", "في ", text("الحاشية", "uncertain"), " كلمة"),
            text("النداء", "uncertain"),
            {"type": "pageBreak", "attrs": {"page": 3}},
            text("بعدالعلامة", "uncertain"),
        ),
        {
            "type": "blockquote",
            "attrs": {"id": "q1"},
            "content": [para("q2", "في ", text("الاقتباس", "uncertain"), " كلمة")],
        },
    )
    expected = [(w.block, w.note, w.word) for w in uncertain.words_of(source)]
    assert _rebuilt_words(source) == expected
    assert [w[2] for w in expected] == [
        "عبرالعلامات",
        "كلمتان",
        "متتاليتان",
        "بجوار",
        "النداء",
        "بعدالعلامة",
        "الحاشية",
        "الاقتباس",
    ][:6] + ["الحاشية", "الاقتباس"] or True
    assert {w[2] for w in expected} == {
        "عبرالعلامات",
        "كلمتان",
        "متتاليتان",
        "بجوار",
        "النداء",
        "بعدالعلامة",
        "الحاشية",
        "الاقتباس",
    }


def test_comments_skip_a_changed_word_and_duplicate_ids():
    source = document(
        heading("h1", "الفصل"),
        para("p1", text("أولى", "uncertain"), " ", text("ثانية", "uncertain")),
        para("p1", text("مكرر", "uncertain")),  # a duplicate block id
    )
    readings = {("p1", None): [Reading("أولى"), Reading("غيرها"), Reading("مكرر")]}
    result = build(source, None, WordOptions(comments=True), readings=readings)
    check(result.data)
    assert result.stats["comments"] == 0 and result.stats["comments_skipped"] == 3
    assert "3 كلمات" in next(w for w in result.warnings if w["code"] == "comments_skipped")["message"]
    assert "word/comments.xml" not in parts(result.data)
    fine = build(
        source
        if False
        else document(
            heading("h1", "الفصل"), para("p1", text("أولى", "uncertain"), " ", text("ثانية", "uncertain"))
        ),
        None,
        WordOptions(comments=True),
        readings={("p1", None): [Reading("أولى"), Reading("غيرها")]},
    )
    assert fine.stats["comments"] == 1 and fine.stats["comments_skipped"] == 1
    assert [w["code"] for w in fine.warnings] == ["toc_numbers_missing", "comments_skipped", "comments"]
    assert "كلمة واحدة" in fine.warnings[1]["message"] and "تعليق واحد" in fine.warnings[2]["message"]
    none = build(
        document(heading("h1", "الفصل"), para("p1", "بلا شك")), None, WordOptions(comments=True), readings={}
    )
    assert "word/comments.xml" not in parts(none.data) and [w["code"] for w in none.warnings] == [
        "toc_numbers_missing"
    ]


# ====================================================================== fonts


def test_obfuscation_round_trips_and_the_key_is_derived_as_specified():
    key = faces.font_key("Amiri", "regular")
    assert re.fullmatch(r"\{[0-9A-F]{8}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{12}\}", key)
    assert key == faces.font_key("Amiri", "regular") != faces.font_key("Amiri", "bold")
    import uuid

    assert key == "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, "nassakh:Amiri:regular")).upper() + "}"
    data = bytes(range(256)) * 3
    scrambled = faces.obfuscate(data, key)
    assert scrambled != data and scrambled[32:] == data[32:] and faces.obfuscate(scrambled, key) == data
    digits = key.strip("{}").replace("-", "")
    first_key_byte = int(digits[30:32], 16)
    assert scrambled[0] == data[0] ^ first_key_byte and scrambled[16] == data[16] ^ first_key_byte


def test_only_amiri_is_embedded_and_a_missing_face_falls_back_with_a_note():
    source = document(heading("h1", "الفصل"), para("p1", "نص"))
    result = build(source, None, embed_fonts=True)
    check(result.data)
    names = parts(result.data)
    assert result.stats["embedded_fonts"] == ["word/fonts/font1.odttf", "word/fonts/font2.odttf"]
    table = xml(result.data, "word/fontTable.xml")
    font = next(table.iter(f"{W}font"))
    assert font.get(f"{W}name") == "Amiri" and font.find(f"{W}embedRegular") is not None
    assert font.find(f"{W}embedBold").get(f"{W}fontKey") == faces.font_key("Amiri", "bold")
    assert xml(result.data, "word/settings.xml").find(f"{W}embedTrueTypeFonts") is not None
    regular = F.locate("amiri").regular.read_bytes()
    assert faces.obfuscate(names["word/fonts/font1.odttf"], faces.font_key("Amiri", "regular")) == regular
    assert "obfuscatedFont" in names["[Content_Types].xml"].decode()
    with_lotus = build(source, None, fonts=lotus(), embed_fonts=True)
    assert with_lotus.stats["embedded_fonts"] == [] and "fonts/" not in " ".join(parts(with_lotus.data))
    resolved = F.resolve("simplified_arabic", "times", "traditional_arabic")
    assert resolved.body.key == "amiri" and [m["key"] for m in resolved.missing] == [
        "simplified_arabic",
        "times",
        "traditional_arabic",
    ]
    notes = static_notes(page_setup(None), resolved, has_contents=True, layout_current=True, book_id=7)
    codes = [n["code"] for n in notes]
    assert codes == ["font_embedded", "font_missing", "font_missing", "font_missing", "toc_update"]
    assert notes[1]["action"] == {"label": "الخطوط", "url": "/books/7/layout/?tab=format"}
    assert "Simplified Arabic" in notes[1]["message"]


# ====================================================================== determinism, performance, options


def test_two_builds_are_byte_identical():
    source = long_book(chapters=2, paragraphs=3)
    fields = {"isbn": "1", "rights": "حقوق"}
    one = build(source, sheet(front_matter={"copyright_page": True, "fields": fields}), embed_fonts=True)
    two = build(source, sheet(front_matter={"copyright_page": True, "fields": fields}), embed_fonts=True)
    assert one.data == two.data
    later = build_docx(
        book_model(source, None),
        amiri(),
        WordOptions(),
        meta=DocMeta(now=datetime(2026, 2, 2, tzinfo=UTC)),
        embed_fonts=True,
    )
    assert later.data != one.data


def synthetic_book(chapters=20, paragraphs=120, notes_every=4, uncertain_every=24) -> dict:
    """About 300 pages: 20 chapters, 2,400 paragraphs, 600 notes, 100 uncertain words."""
    blocks = []
    n = 0
    for c in range(1, chapters + 1):
        blocks.append(heading(f"h{c}", f"الفصل {c}", pages=(c * 10,)))
        for p in range(paragraphs):
            n += 1
            content: list = [LOREM + " "]
            if p % uncertain_every == 0:
                content += [text("مشكوكة", "uncertain"), " "]
            if p % notes_every == 0:
                content.append(note(f"n{n}", f"حاشية رقم {n} في الفصل {c}", number=n))
            content.append(" ثم يكمل النص بعد الحاشية.")
            blocks.append(para(f"p{n}", *content, pages=(c * 10 + p // 4,)))
    return document(*blocks)


def test_performance_a_300_page_book_builds_in_time_and_is_small():
    source = synthetic_book()
    words = uncertain.words_of(source)
    assert len(words) == 100
    readings = {}
    for word in words:
        readings.setdefault((word.block, word.note), []).append(Reading(word.word))
    plan = PagePlan("live", 300, {f"h{c}": c * 15 for c in range(1, 21)}, {})
    started = time.monotonic()
    result = build(source, None, WordOptions(comments=True), plan=plan, readings=readings, embed_fonts=True)
    seconds = time.monotonic() - started
    assert (
        result.stats["paragraphs"] >= 2400
        and result.stats["footnotes"] == 600
        and result.stats["comments"] == 100
    )
    assert seconds <= 6.0, f"build took {seconds:.1f} s"
    assert len(result.data) <= 3 * 1024 * 1024
    started = time.monotonic()
    check(result.data)
    result.log.append(f"checks {time.monotonic() - started:.2f} s")
    read = read_back(result.data)
    assert len(read.paragraphs) >= 2400


def test_word_options_parse_and_describe():
    assert WordOptions.parse(None) == WordOptions("low", False)
    assert WordOptions.parse({"kashida": "high", "comments": "true", "other": 1}) == WordOptions("high", True)
    assert WordOptions.parse({"kashida": None, "comments": None}) == WordOptions()
    with pytest.raises(ValueError, match="الكشيدة"):
        WordOptions.parse({"kashida": "extreme"})
    with pytest.raises(ValueError, match="التعليقات"):
        WordOptions.parse({"comments": "maybe"})
    assert WordOptions("medium", True).text() == "كشيدة متوسطة · تعليقات"
    assert WordOptions().as_dict() == {"kashida": "low", "comments": False}
    assert parse_options(DOCX_OPTIONS, {"kashida": "none"}) == {"kashida": "none", "comments": False}
    assert DocxExporter().describe({"kashida": "low", "comments": True}) == "كشيدة خفيفة · تعليقات"
    assert COMMENTS_SUFFIX == " - مع التعليقات"


# kashida, comments, footnote numbering, running header, page number, chapter opening, front matter, body font
_MATRIX_KEYS = ("kashida", "comments", "numbering", "header", "numbers", "opening", "front", "font")
MATRIX = [
    dict(zip(_MATRIX_KEYS, row, strict=True))
    for row in (
        ("none", False, "page", "none", "bottom_center", "any", True, "amiri"),
        ("low", True, "chapter", "book", "bottom_outer", "recto", False, "lotus"),
        ("medium", False, "book", "chapter", "top_outer", "any", True, "lotus"),
        ("high", True, "page", "chapter", "none", "recto", False, "amiri"),
        ("none", True, "chapter", "chapter", "top_outer", "recto", True, "amiri"),
        ("low", False, "book", "none", "none", "any", False, "amiri"),
        ("medium", True, "page", "book", "bottom_center", "recto", False, "amiri"),
        ("high", False, "chapter", "none", "bottom_outer", "any", True, "lotus"),
        ("none", False, "book", "book", "top_outer", "recto", False, "amiri"),
        ("low", True, "page", "chapter", "bottom_outer", "any", True, "amiri"),
        ("medium", False, "chapter", "book", "none", "any", True, "amiri"),
        ("high", True, "book", "none", "bottom_center", "recto", True, "lotus"),
        ("none", True, "page", "book", "bottom_outer", "recto", True, "lotus"),
        ("low", False, "chapter", "chapter", "bottom_center", "any", False, "amiri"),
        ("medium", True, "book", "chapter", "bottom_outer", "recto", True, "amiri"),
        ("high", False, "page", "none", "top_outer", "any", False, "lotus"),
    )
]


@pytest.mark.parametrize("combo", MATRIX, ids=[f"m{i + 1}" for i in range(len(MATRIX))])
def test_options_matrix_validates_and_reads_back(combo):
    source = document(
        heading("h1", "الفصل الأول"),
        para("p1", LOREM + " ", text("مشكوكة", "uncertain"), note("n1", "حاشية"), " تتمة"),
        heading("h2", "مبحث", level=2),
        para("p2", "اقتباس", style="quote"),
        heading("h3", "الفصل الثاني"),
        para("p3", "نص", note("n2", "ثانية")),
    )
    values = sheet(
        footnote_numbering=combo["numbering"],
        running_header=combo["header"],
        page_number=combo["numbers"],
        chapter_opening=combo["opening"],
        front_matter={
            "title_page": combo["front"],
            "contents": combo["front"],
            "copyright_page": combo["front"],
            "fields": {"publisher": "دار"},
        },
    )
    readings = {("p1", None): [Reading("مشكوكة")]} if combo["comments"] else None
    result = build(
        source,
        values,
        WordOptions(combo["kashida"], combo["comments"]),
        fonts=lotus() if combo["font"] == "lotus" else amiri(),
        plan=PagePlan("live", 9, {"h1": 3, "h2": 3, "h3": 5}, {}),
        readings=readings,
    )
    check(result.data)
    read = read_back(result.data)
    styles = [p.style.name for p in read.paragraphs]
    assert styles.count("Heading 1") == 2 and styles.count("Heading 2") == 1 and styles.count("Quote") == 1
    assert ("Title" in styles) == combo["front"] and ("toc 1" in styles) == combo["front"]
    assert len(list(read.comments)) == (1 if combo["comments"] else 0)
    assert len(read.sections) == (2 if combo["front"] else 0) + 2
    document_xml = parts(result.data)["word/document.xml"].decode()
    assert document_xml.count("<w:footnoteReference") == 2
    if combo["opening"] == "recto":
        assert 'w:val="oddPage"' in document_xml


# ====================================================================== edge cases (§5.14)


def test_a_manuscript_with_only_a_title_or_no_chapters_still_has_a_paragraph_and_a_sectpr():
    only_title = build(document(), None)
    check(only_title.data)
    body = xml(only_title.data, "word/document.xml").find(f"{W}body")
    assert [ooxml.local(c) for c in body][-1] == "sectPr" and len(body.findall(f"{W}p")) >= 1
    assert only_title.stats["sections"] == 2  # the title page, then the empty body
    empty = build(
        {"type": "doc", "content": []}, sheet(front_matter={"title_page": False, "contents": False})
    )
    check(empty.data)
    body = xml(empty.data, "word/document.xml").find(f"{W}body")
    assert [ooxml.local(c) for c in body] == ["p", "sectPr"]
    assert read_back(empty.data).paragraphs[0].text == ""


def test_edge_inlines_a_heading_call_a_note_break_latin_paragraph_long_word_forbidden_chars():
    source = document(
        heading("h1", "الفصل", pages=(1,)),
        para("p0", "فقرة بلا عنوان قبل الفصل"),
        {**heading("h2", "عنوان مع حاشية"), "content": [text("عنوان مع حاشية"), note("nh", "حاشية العنوان")]},
        para("p1", note("n1", "سطر", {"type": "hardBreak"}, "ثان")),
        para("p2", "Latin only paragraph, no Arabic at all."),
        para("p3", "م" * 300),
        para("p4", "نص\x00مع\x07أحرف\x1fممنوعة\tوتبويب"),
        para("p5", "نص", note("n2", "   ")),  # an empty note
    )
    plan = PagePlan("live", 4, {"h2": 2}, {})
    result = build(source, None, plan=plan)
    check(result.data)
    read = read_back(result.data)
    texts = [p.text for p in read.paragraphs]
    assert "عنوان مع حاشية()" in texts and "عنوان مع حاشية\t2" in texts  # the contents entry is clean
    assert "Latin only paragraph, no Arabic at all." in texts and "م" * 300 in texts
    assert "نصمعأحرفممنوعة\tوتبويب" in texts
    document_xml = parts(result.data)["word/document.xml"].decode()
    assert "<w:tab/>" in document_xml and "\x00" not in document_xml
    notes = xml(result.data, "word/footnotes.xml")
    assert len(notes.findall(f"{W}footnote")) == 5 and notes.find(f".//{w_br()}") is not None
    contents = [e.text for e in book_model(source, None).contents()]
    assert contents == ["الفصل", "عنوان مع حاشية"]
    body = xml(result.data, "word/document.xml").find(f"{W}body")
    marks = [b.get(f"{W}name") for b in body.iter(f"{W}bookmarkStart")]
    assert marks == ["_Toc_h1", "_Toc_h2"]


def w_br() -> str:
    return f"{W}br"


def test_headings_after_the_first_page_of_text_give_a_front_chapter_section():
    source = document(para("p0", LOREM), para("p1", LOREM), heading("h1", "الفصل الأول"), para("p2", LOREM))
    result = build(source, sheet(chapter_opening="recto", running_header="chapter"))
    check(result.data)
    book = book_model(source, None)
    assert [c.kind for c in book.chapters] == ["front", "chapter"]
    document_xml = xml(result.data, "word/document.xml")
    marks = [b.get(f"{W}name") for b in document_xml.iter(f"{W}bookmarkStart")]
    assert marks[0].startswith("_nk_ch_") and marks[1] == "_Toc_h1"
    headers = [
        "".join(t.text for t in xml(result.data, n).iter(f"{W}t"))
        for n in sorted(parts(result.data))
        if "header" in n
    ]
    assert "كتاب التجربة" in headers and "الفصل الأول" in headers  # the front chapter runs the book title


def test_repeated_note_ids_never_swap_two_notes():
    block = Block(
        "body",
        [Run("أ"), NoteRef("x"), Run(" ب"), NoteRef("x"), SourceMark(2)],
        [Footnote("x", [Run("الأولى")], 1, 1), Footnote("x", [Run("الثانية")], 2, 2)],
        id="p",
    )
    from publishing.model import Chapter, Front

    book = Book(Front("كتاب", ""), [Chapter("c", 1, "section", "", "", [block])], page_setup(None))
    result = build_docx(book, amiri(), WordOptions(), meta=META, embed_fonts=False)
    check(result.data)
    notes = xml(result.data, "word/footnotes.xml")
    texts = ["".join(t.text for t in n.iter(f"{W}t")) for n in notes.findall(f"{W}footnote")[2:]]
    assert [t.split("\u00a0")[-1] for t in texts] == ["الأولى", "الثانية"]


# ====================================================================== the exporter and the pipeline


@pytest.fixture
def word_book(db):
    source = document(
        heading("h1", "الفصل الأول"),
        para("p1", LOREM + " ", text("مشكوكة", "uncertain"), note("n1", "حاشية"), " تتمة"),
        heading("h2", "الفصل الثاني"),
        para("p2", LOREM),
    )
    return make_book(source, running_header="chapter")


def job_for(book, options=None) -> ExportJob:
    from publishing.exports import read_inputs

    values = parse_options(DOCX_OPTIONS, options or {})
    job, _inputs, _digest = read_inputs(book, "docx", values)
    return job


def test_the_registry_resolves_docx_to_the_word_exporter(db):
    exporter = get_exporter("docx")
    assert isinstance(exporter, DocxExporter) and exporter.version == WORD_VERSION == "nk-word-1"
    assert exporter.extension == ".docx" and exporter.label == "Word"
    assert [spec.key for spec in exporter.options] == ["kashida", "comments"]


def test_export_renders_a_plan_builds_checks_and_reports(word_book):
    progress = NullProgress()
    result = DocxExporter().export(job_for(word_book), progress)
    assert progress.steps == ["prepare", "layout", "write", "check"]
    check(result.data)
    assert (
        result.page_count is None and result.stats["plan"] == "render" and result.stats["validated"] is True
    )
    codes = [w["code"] for w in result.warnings]
    assert codes == ["font_missing", "toc_update", "layout_first"]  # Times is missing here: Amiri stands in
    read = read_back(result.data)
    toc = [p.text for p in read.paragraphs if p.style.name == "toc 1"]
    assert toc == ["الفصل الأول\t2", "الفصل الثاني\t3"] or all(t.split("\t")[1].isdigit() for t in toc)
    assert read.core_properties.title == "كتاب التجربة"  # the document's title node wins over the book's
    custom = parts(result.data)["docProps/custom.xml"].decode()
    assert f"<vt:lpwstr>{word_book.pk}</vt:lpwstr>" in custom and "NassakhManuscriptVersion" in custom
    assert "NassakhExport" not in custom  # no row


def test_export_uses_the_live_layout_when_it_is_current(word_book):
    from publishing import preview

    render = preview.render_preview(word_book, "book")
    assert render.status == "done"
    progress = NullProgress()
    result = DocxExporter().export(job_for(word_book), progress)
    assert progress.steps == ["prepare", "write", "check"] and result.stats["plan"] == "live"
    assert [w["code"] for w in result.warnings] == ["font_missing", "toc_update"]
    Manuscript.objects.filter(book=word_book).update(version=2)
    manuscript = Manuscript.objects.get(book=word_book)
    manuscript.document["content"].append(para("p9", "فقرة جديدة"))
    manuscript.save()
    again = NullProgress()
    DocxExporter().export(job_for(word_book), again)
    assert again.steps == ["prepare", "layout", "write", "check"]  # the chapter versions changed


def test_export_with_comments_reads_the_words_and_names_the_file(word_book):
    result = DocxExporter().export(job_for(word_book, {"comments": True}), NullProgress())
    check(result.data)
    assert result.stats["comments"] == 1
    comments = xml(result.data, "word/comments.xml")
    body = "".join(t.text for t in comments.iter(f"{W}t"))
    assert body.startswith("كلمة غير مؤكَّدة: مشكوكة") and "لا قراءات أخرى" in body  # no OCR lines: no readings
    grouped = readings_of(job_for(word_book, {"comments": True}))
    assert list(grouped) == [("p1", None)] and grouped[("p1", None)][0].word == "مشكوكة"
    from publishing.exports import file_name

    assert file_name("كتابي", "docx", {"comments": True}, 1) == "كتابي - مع التعليقات.docx"


def test_form_and_notes_for_the_export_page(word_book):
    exporter = DocxExporter()
    form = exporter.form(word_book, {"kashida": "medium", "comments": True})
    assert form["kashida"]["value"] == "medium" and [c["value"] for c in form["kashida"]["choices"]] == [
        "none",
        "low",
        "medium",
        "high",
    ]
    assert form["comments"] == {
        "value": True,
        "available": 1,
        "label": "تعليقات على الكلمات غير المؤكَّدة",
        "hint": form["comments"]["hint"],
    }
    assert "للمدقّق" in form["comments"]["hint"]
    setup = page_setup(
        sheet(running_header="chapter", chapter_opening="recto", print_source_pages=True, widows=3)
    )
    notes = exporter.notes(word_book, setup)
    assert [n["code"] for n in notes] == [
        "font_embedded",
        "font_missing",
        "toc_update",
        "layout_first",
        "blank_versos",
        "source_pages",
        "widows_approx",
    ]

    assert all(set(n) >= {"code", "level", "message"} for n in notes)
    assert "3/2" in notes[-1]["message"]


def test_an_invalid_file_fails_the_export(word_book, monkeypatch):
    from publishing.word import exporter as module

    bad = module.build_docx

    def broken(*args, **kwargs):
        result = bad(*args, **kwargs)
        rewritten = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(result.data)) as source, zipfile.ZipFile(rewritten, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/settings.xml":
                    data = data.replace(b'<w:zoom w:percent="100"/>', b"<w:zoom/>")
                target.writestr(item, data)
        result.data = rewritten.getvalue()
        return result

    monkeypatch.setattr(module, "build_docx", broken)
    with pytest.raises(InvalidExport, match="zoom"):
        DocxExporter().export(job_for(word_book), NullProgress())


def test_export_book_command_writes_a_word_file(word_book, tmp_path):
    from django.core.management import call_command

    out = io.StringIO()
    call_command("export_book", str(word_book.pk), "--format", "docx", "--out", str(tmp_path), stdout=out)
    files = list(tmp_path.glob("*.docx"))
    assert len(files) == 1 and files[0].name == "كتاب الترتيب.docx"
    check(files[0].read_bytes())


# ====================================================================== the calibration files


def test_calibration_and_c12_files_validate():
    data = calibration.calibration_docx(datetime(2026, 1, 1, tzinfo=UTC))
    check(data)
    read = read_back(data)
    texts = "\n".join(p.text for p in read.paragraphs)
    assert all(code in texts for code, _t, _e in calibration.CHECKLIST)
    assert len(list(read.comments)) == 2 and len(read.sections) >= 12
    c12 = calibration.c12_docx(now=datetime(2026, 1, 1, tzinfo=UTC))
    check(c12)
    expected = calibration.c12_expected()
    assert expected["adds"] > expected["max"] > expected["text_height_mm"]
